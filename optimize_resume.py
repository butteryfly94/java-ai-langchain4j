from docx import Document

SOURCE = r"D:\bak\ai-langchain4j-prd\马亚平_优化版.docx"


def replace_paragraph(paragraph, text):
    """Preserve the paragraph's layout while replacing its visible text."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


doc = Document(SOURCE)

replacements = {
    "3. 具备RAG检索增强系统的完整设计与落地能力：基于 LangChain4j构建离线/在线双链路架构——离线父子分块索引管线（子块检索保精度、父块上下文保完整），在线“查询扩展 → kNN+BM25双路混合检索 → RRF 融合 → Rerank 重排”完整管线，提升检索准确率。\t":
        "3. 具备 RAG 检索增强系统的完整设计与落地能力：基于 LangChain4j 构建离线索引与在线问答双链路；采用父子分块、查询扩展、kNN + BM25 混合检索、RRF 融合和 Rerank 重排，并实现固定 RAG 与 Agentic RAG 两种问答模式。",
    "• viber coding“警小宁”app，基于 Flutter 开发无人巡逻车双模式（车载/手机）移动应用集成大模型流式问答、问路引导与高德地图导航，实现巡逻任务管理与实现云台、探照灯、警灯、无人机机盖等车辆远程控制。":
        "• 采用 Vibe Coding 开发“警小宁”Flutter 移动应用，支持车载/手机双模式；集成大模型流式问答、问路引导和高德地图导航，并实现巡逻任务管理及云台、探照灯、警灯、无人机机盖等车辆远程控制。",
    "•  独立设计并落地离线索引管线：实现父子分块（parent-child chunking）两级切分策略——文档切分为 1000 字父块存原文索引、300字子块向量化存检索索引，子块携带 parentId 元数据实现“子块检索保召回精度、父块回查保上下文完整”，解决单一分块粒度下“检索准则答案残缺、上下文全则召回下降”的固有矛盾。":
        "• 独立设计并落地离线知识库索引：采用父子分块策略，将文档切为约 1000 字父块和约 300 字子块；父块保存完整上下文，子块向量化并携带 parentId 用于精准召回后回查原文，兼顾召回精度与上下文完整性。",
    "  • 构建在线混合检索管线：设计“ExpandingQueryTransformer 查询扩展 → kNN 向量检索 + BM25 全文检索双路并发召回 → 应用层RRF 融合（k=60）→ 父块扩展 → gte-rerank 交叉编码器重排 → top-3 注入”完整链路；因 ES 服务端 RRF需企业许可证，自行实现应用层 RRF 融合算法，零商业依赖达成同等效果。":
        "• 构建固定 RAG 在线检索链路：查询扩展后并行执行 Elasticsearch kNN 向量检索与 BM25 全文检索，对子块结果进行应用层 RRF 融合（k=60）、父块回查和 gte-rerank 交叉编码器重排，最终将 Top-3 上下文注入模型；规避 ES 服务端 RRF 的企业版依赖。",
    "  • 实现 Agentic RAG 与 MCP 协议接入：将知识库检索封装为 Tool 由 LLM自主调度检索时机与查询改写（与固定管线形成两种模式对比），并通过 MCP 协议（stdio 传输）接入 filesystemserver，实现本地 @Tool 与 MCP 远程工具的统一编排。":
        "• 实现 Agentic RAG 与 MCP 接入：将 ES 向量检索封装为本地 @Tool，由 LLM 按需决定检索时机、查询改写与多次检索；通过 MCP stdio 接入 filesystem server，在知识库白名单目录内提供文件列举和原文读取能力，实现本地工具与 MCP 工具的统一编排，并与固定 RAG 形成分层路由模式。",
    "  • 建立检索质量评测体系：自建问答评测集，编写 Recall 评测用例量化验证检索链路效果（混合检索 + RRF + Rerank 完整链路Recall 达【补充实测数据】），支持 rerank 开关对比实验，以数据驱动检索参数（分块大小、RRF k 值、top-N）调优。":
        "• 建立检索质量评测基础设施：自建“问题—目标文档”评测集，直接调用 RetrievalAugmentor 对查询扩展、混合检索、RRF、父块回查与 Rerank 全链路计算 Recall@3；支持 Rerank 开关对比，并据此调优分块大小、RRF k 值和 Top-N 参数。",
    "  •工程化细节：全链路降级设计（rerank 失败降级 RRF 顺序、父块索引未命中降级返回子块，保证服务可用性）；DashScope向量接口批量上限的分批处理；索引全量重建幂等性设计；基于 Docker Compose 的 ES 一键环境搭建。":
        "• 工程化保障：实现 Rerank 失败降级为 RRF 排序、父块未命中降级返回子块；针对 DashScope 向量接口批量上限进行分批处理，保证索引全量重建幂等，并通过 Docker Compose 一键启动 Elasticsearch 环境。",
}

unmatched = set(replacements)
for paragraph in doc.paragraphs:
    original = paragraph.text
    if original in replacements:
        replace_paragraph(paragraph, replacements[original])
        unmatched.remove(original)

if unmatched:
    raise ValueError(f"Unmatched paragraphs: {unmatched}")

doc.core_properties.comments = "优化版：城安学堂聊天机器人项目描述已按实际实现更新。"
doc.save(SOURCE)
